import os
import re
from docx import Document
from docx.shared import Pt

MD_PATH = r'd:\Shinnovation\leox\Akash_Sanjay_More_Resume.md'
DOCX_PATH = r'd:\Shinnovation\leox\Akash_Sanjay_More_Resume.docx'


def add_heading(doc, text, level):
    # level: 0 for Title, 1 for Heading 1, etc.
    doc.add_heading(text, level=level)


def add_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    # Simple markdown inline formatting: **bold**, *italic*
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.italic = True
        else:
            p.add_run(part)
    return p


def main():
    if not os.path.exists(MD_PATH):
        print(f"Markdown file not found: {MD_PATH}")
        return
    doc = Document()
    with open(MD_PATH, encoding='utf-8') as f:
        for raw_line in f:
            line = raw_line.rstrip('\n')
            if not line.strip():
                continue
            if line.startswith('# '):
                add_heading(doc, line[2:].strip(), level=0)
            elif line.startswith('## '):
                add_heading(doc, line[3:].strip(), level=1)
            elif line.startswith('### '):
                add_heading(doc, line[4:].strip(), level=2)
            elif line.startswith('- '):
                add_paragraph(doc, line[2:].strip(), style='List Bullet')
            else:
                add_paragraph(doc, line.strip())
    doc.save(DOCX_PATH)
    print(f"Docx generated at {DOCX_PATH}")

if __name__ == '__main__':
    main()
